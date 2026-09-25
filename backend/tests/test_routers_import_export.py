"""Tests du router `import-export`.

`test_robustesse_sauvegardes.py` couvre déjà en profondeur `/export/complet`
et `/import` (zip-slip, rotation, transactionnalité, confirmation). Ce
fichier couvre ce qui restait non exercé via HTTP : `/export` (classeur
XLSX simple), `/sauvegardes` (liste), et un aller-retour export→import
minimal en conditions normales."""
import io
import zipfile
from datetime import date

import openpyxl
import pytest

import models


@pytest.fixture()
def classe(db_session):
    c = models.Classes(niveau="7ème Année", nom="7ème Année A", frais_inscription=10000, mensualite=5000)
    db_session.add(c)
    db_session.commit()
    db_session.refresh(c)
    return c


class TestExportSimple:
    def test_export_xlsx_valide(self, client, auth_headers, classe):
        resp = client.get("/api/import-export/export", headers=auth_headers)
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]
        assert "collegeaureole_export_" in resp.headers["content-disposition"]

        classeur = openpyxl.load_workbook(io.BytesIO(resp.content), read_only=True)
        assert "classes" in classeur.sheetnames
        assert "eleves" in classeur.sheetnames
        feuille_classes = classeur["classes"]
        lignes = list(feuille_classes.iter_rows(values_only=True))
        entetes = lignes[0]
        assert "nom" in entetes
        valeurs = [dict(zip(entetes, ligne)) for ligne in lignes[1:]]
        assert any(v["nom"] == "7ème Année A" for v in valeurs)

    def test_export_ne_contient_pas_le_blob_documents(self, client, auth_headers):
        resp = client.get("/api/import-export/export", headers=auth_headers)
        classeur = openpyxl.load_workbook(io.BytesIO(resp.content), read_only=True)
        entetes = next(classeur["documents"].iter_rows(values_only=True))
        assert "contenu" not in entetes

    def test_export_sans_authentification_401(self, client):
        resp = client.get("/api/import-export/export")
        assert resp.status_code == 401


class TestListeSauvegardes:
    def test_liste_vide_par_defaut(self, client, auth_headers, tmp_path, monkeypatch):
        monkeypatch.setenv("AUREOLE_SAUVEGARDES_DIR", str(tmp_path))
        resp = client.get("/api/import-export/sauvegardes", headers=auth_headers)
        assert resp.status_code == 200
        assert resp.json() == {"sauvegardes": []}

    def test_liste_apres_creation(self, client, auth_headers, tmp_path, monkeypatch):
        monkeypatch.setenv("AUREOLE_SAUVEGARDES_DIR", str(tmp_path))
        creation = client.post("/api/import-export/sauvegarde", headers=auth_headers)
        assert creation.status_code == 200, creation.text
        nom = creation.json()["sauvegarde"]

        liste = client.get("/api/import-export/sauvegardes", headers=auth_headers).json()["sauvegardes"]
        noms = [s["nom"] for s in liste]
        assert nom in noms

    def test_sans_authentification_401(self, client):
        resp = client.get("/api/import-export/sauvegardes")
        assert resp.status_code == 401


class TestExportImportAllerRetour:
    def test_restauration_conserve_les_donnees(self, client, auth_headers, classe, db_session):
        archive = client.get("/api/import-export/export/complet", headers=auth_headers).content
        assert zipfile.is_zipfile(io.BytesIO(archive))

        # On modifie l'état courant après l'export pour vérifier que
        # l'import restaure bien l'état capturé (et pas l'état courant).
        db_session.query(models.Classes).delete()
        db_session.commit()
        assert db_session.query(models.Classes).count() == 0

        resp = client.post(
            "/api/import-export/import",
            data={"remplacer": "true"},
            files={"fichier": ("sauvegarde.zip", io.BytesIO(archive), "application/zip")},
            headers={**auth_headers, "X-Confirm": "RESTAURATION-DONNEES"},
        )
        assert resp.status_code == 200, resp.text
        assert resp.json()["lignes_importees"] > 0

        db_session.expire_all()
        restaurees = db_session.query(models.Classes).all()
        assert any(c.nom == "7ème Année A" for c in restaurees)

    def test_fichier_vide_refuse(self, client, auth_headers):
        resp = client.post(
            "/api/import-export/import",
            data={"remplacer": "true"},
            files={"fichier": ("vide.zip", io.BytesIO(b""), "application/zip")},
            headers={**auth_headers, "X-Confirm": "RESTAURATION-DONNEES"},
        )
        assert resp.status_code == 400
        assert "vide" in resp.json()["detail"].lower()

    def test_format_invalide_refuse(self, client, auth_headers):
        resp = client.post(
            "/api/import-export/import",
            data={"remplacer": "true"},
            files={"fichier": ("pas_un_zip.zip", io.BytesIO(b"ceci n'est pas un zip"), "application/zip")},
            headers={**auth_headers, "X-Confirm": "RESTAURATION-DONNEES"},
        )
        assert resp.status_code == 400
